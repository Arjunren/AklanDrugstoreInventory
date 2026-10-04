from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "docx" / "Arjunren_Valdez_IntermediateMobileProg_04TaskPerformance1.docx"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)

BLACK = "000000"
DARK_BLUE = "17365D"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F5F5F5"
BORDER_GRAY = "D9D9D9"
GREEN = "167A53"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER_GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_bottom_border(paragraph, color=BORDER_GRAY):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(document, headers, rows, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for index, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[index]
        cell.width = width
        set_cell_shading(cell, DARK_BLUE)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)

    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        for column_index, (value, width) in enumerate(zip(row_values, widths)):
            cell = cells[column_index]
            cell.width = width
            if row_index % 2 == 1:
                set_cell_shading(cell, LIGHT_BLUE)
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index in (0, len(headers) - 1) else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(str(value))
            run.font.size = Pt(9)
    return table


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.font.size = Pt(11)
    return paragraph


def add_code_file(document, relative_path, title, start_new_page=True):
    heading = document.add_heading(title, level=1)
    heading.paragraph_format.page_break_before = start_new_page
    heading.paragraph_format.keep_with_next = True
    path_line = document.add_paragraph()
    path_line.paragraph_format.space_after = Pt(6)
    run = path_line.add_run("Source file  " + relative_path)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(89, 89, 89)

    code = (ROOT / relative_path).read_text(encoding="utf-8").replace("\t", "    ")
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.08)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.widow_control = False
    shade_paragraph(paragraph, LIGHT_GRAY)
    run = paragraph.add_run(code)
    run.font.name = "Consolas"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
    run.font.size = Pt(6.5)
    run.font.color.rgb = RGBColor(38, 38, 38)


document = Document()
section = document.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = document.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title_style.font.size = Pt(28)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_p_pr = title_style.element.get_or_add_pPr()
title_border = title_p_pr.find(qn("w:pBdr"))
if title_border is not None:
    title_p_pr.remove(title_border)

for style_name, size in (("Heading 1", 18), ("Heading 2", 13)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(6)

document.core_properties.title = "Aklan Drugstore Inventory Project Documentation"
document.core_properties.subject = "Intermediate Mobile Programming Task Performance"
document.core_properties.author = "Arjunren Valdez"
document.core_properties.comments = "Source code and project documentation"

footer = section.footer
footer_paragraph = footer.paragraphs[0]
footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer_paragraph.add_run("Aklan Drugstore Inventory  |  Arjunren Valdez")
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor(102, 102, 102)

# Cover page
document.add_paragraph()
document.add_paragraph()
eyebrow = document.add_paragraph()
eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
eyebrow_run = eyebrow.add_run("INTERMEDIATE MOBILE PROGRAMMING")
eyebrow_run.bold = True
eyebrow_run.font.size = Pt(11)
eyebrow_run.font.color.rgb = RGBColor(89, 89, 89)

title = document.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Aklan Drugstore Inventory Project Documentation")

subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle_run = subtitle.add_run("Task Performance 1")
subtitle_run.font.size = Pt(15)
subtitle_run.font.color.rgb = RGBColor(89, 89, 89)

document.add_paragraph()
cover_rows = [
    ("Developer", "Arjunren Valdez"),
    ("Technology", ".NET MAUI and C#"),
    ("Platforms", "Windows and Android"),
    ("Project type", "Simple local student application"),
    ("Date", "October 4, 2026"),
]
cover = document.add_table(rows=0, cols=2)
cover.alignment = WD_TABLE_ALIGNMENT.CENTER
cover.autofit = False
for label, value in cover_rows:
    cells = cover.add_row().cells
    cells[0].width = Inches(1.55)
    cells[1].width = Inches(3.65)
    set_cell_shading(cells[0], LIGHT_BLUE)
    for cell in cells:
        set_cell_borders(cell)
        set_cell_margins(cell, top=120, bottom=120)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    label_run = cells[0].paragraphs[0].add_run(label)
    label_run.bold = True
    label_run.font.size = Pt(10)
    value_run = cells[1].paragraphs[0].add_run(value)
    value_run.font.size = Pt(10)

credit = document.add_paragraph()
credit.alignment = WD_ALIGN_PARAGRAPH.CENTER
credit.paragraph_format.space_before = Pt(26)
credit_run = credit.add_run("Developed by Arjunren Valdez")
credit_run.bold = True
credit_run.font.size = Pt(12)

document.add_page_break()

# Project overview
document.add_heading("Project Overview", level=1)
document.add_paragraph(
    "This document contains the project explanation, opening instructions, feature checklist, "
    "and core source code for the Aklan Drugstore Inventory application. The application is a "
    "simple .NET MAUI student project for a local drugstore in Aklan, Panay Island. It replaces "
    "basic manual inventory checking with an offline system that runs on Windows and Android."
)
document.add_paragraph(
    "The project is ready to open using the traditional Visual Studio solution file named "
    "AklanDrugstoreInventory.sln. Product and audit data are stored in in-memory list collections, "
    "so the classroom demonstration does not require SQL, a database server, or internet access. "
    "The list resets when the application closes."
)

document.add_heading("System Login", level=2)
login = document.add_paragraph()
login.add_run("Username  ").bold = True
login.add_run("admin")
login.add_run("\nPassword  ").bold = True
login.add_run("admin123")
document.add_paragraph(
    "Enter the username and password exactly as shown, then select Log In. The fixed account is "
    "only for a classroom demonstration and should not be used as a production security design."
)

document.add_heading("Feature Checklist", level=2)
feature_rows = [
    ("Login and logout", "Local administrator account and validation", "Complete"),
    ("Menu screen", "Dashboard, Inventory, and Audit Trail navigation", "Complete"),
    ("Inventory management", "Add, update, delete, select, and display products", "Complete"),
    ("Data visualization", "Product count, total units, low stock, and bar chart", "Complete"),
    ("Audit trail", "Login, logout, add, update, and delete records", "Complete"),
    ("List storage", "In-memory inventory and audit collections for the open session", "Complete"),
    ("Category suggestions", "Filters previously used categories while the user types", "Complete"),
]
add_table(document, ["Requirement", "Implementation", "Status"], feature_rows, [Inches(1.55), Inches(3.85), Inches(1.0)])

document.add_page_break()
document.add_heading("Open the Project in Visual Studio", level=1)
steps = [
    "Open File Explorer and go to the Inventory project folder.",
    "Double-click AklanDrugstoreInventory.sln. Use the .sln file, not a source file or the output folder.",
    "Wait for Visual Studio to restore the .NET MAUI project packages.",
    "At the top of Visual Studio, choose Windows Machine or an Android emulator.",
    "Press F5 or select the green Start button.",
    "When the login page appears, enter admin and admin123.",
]
for step in steps:
    add_bullet(document, step)

document.add_heading("If Visual Studio Does Not Load the Project", level=2)
document.add_paragraph(
    "Open Visual Studio Installer, select Modify for Visual Studio 2026, and confirm that the "
    ".NET Multi-platform App UI development workload is installed. Also confirm that the .NET 10 "
    "SDK is available. After installation, reopen AklanDrugstoreInventory.sln and allow package "
    "restore to finish before starting the app."
)

document.add_heading("Build Commands", level=2)
commands = (
    "dotnet build AklanDrugstoreInventory.csproj -f net10.0-windows10.0.19041.0\n"
    "dotnet build AklanDrugstoreInventory.csproj -f net10.0-android"
)
command_paragraph = document.add_paragraph()
shade_paragraph(command_paragraph, LIGHT_GRAY)
command_run = command_paragraph.add_run(commands)
command_run.font.name = "Consolas"
command_run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
command_run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
command_run.font.size = Pt(9)

document.add_heading("Project Structure", level=1)
structure_rows = [
    ("AklanDrugstoreInventory.sln", "Traditional Visual Studio solution file"),
    ("AklanDrugstoreInventory.csproj", ".NET MAUI project settings and target platforms"),
    ("MainPage.xaml", "Login, dashboard, inventory, and audit user interface"),
    ("MainPage.xaml.cs", "Login, navigation, CRUD, category suggestions, list data, and audit logic"),
    ("Models", "InventoryItem and AuditEntry data classes"),
    ("Resources", "App icons, splash screen, fonts, styles, and raw assets"),
]
add_table(document, ["File or Folder", "Purpose"], structure_rows, [Inches(2.55), Inches(3.85)])

document.add_page_break()
document.add_heading("Core Source Code", level=1)
document.add_paragraph(
    "The following pages contain the core XAML and C# source used by the project. Standard "
    "platform manifests, generated build files, fonts, icons, and style resources remain in the "
    "project repository."
)

source_files = [
    ("AklanDrugstoreInventory.csproj", "Project File"),
    ("App.xaml", "Application Resources"),
    ("App.xaml.cs", "Application Startup"),
    ("AppShell.xaml", "Application Shell"),
    ("AppShell.xaml.cs", "Application Shell Code"),
    ("MauiProgram.cs", "MAUI Configuration"),
    ("Models/InventoryItem.cs", "Inventory Item Model"),
    ("Models/AuditEntry.cs", "Audit Entry Model"),
    ("MainPage.xaml", "Main User Interface"),
    ("MainPage.xaml.cs", "Main Application Logic"),
]
for index, (relative_path, source_title) in enumerate(source_files):
    add_code_file(document, relative_path, source_title, start_new_page=index > 0)

document.add_page_break()
document.add_heading("Credits and License", level=1)
document.add_paragraph("Designed and developed by Arjunren Valdez.")
document.add_paragraph(
    "The project is released under the MIT License. The repository contains the application "
    "source code, README, license, project assets, solution files, and this editable Word document."
)

document.save(OUTPUT)
print(OUTPUT)
